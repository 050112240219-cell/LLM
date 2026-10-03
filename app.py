"""Goal-based Personal Finance Planning Engine (Flask + SQLite + Gemini via Google AI Studio)."""
import os, io, json, sqlite3, datetime as dt
import requests
from flask import Flask, request, jsonify, send_file, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

app = Flask(__name__)
DB = os.environ.get("DB_PATH", "data/clients.db")
ADMIN_PW = os.environ.get("ADMIN_PASSWORD", "admin123")  # CHANGE in production
DEFAULT_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.6-flash")
os.makedirs(os.path.dirname(DB) or ".", exist_ok=True)


def db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute("create table if not exists clients(id integer primary key autoincrement, created text, name text, payload text)")
    return c


def vnd(x):
    return f"{x:,.0f} VND"


def num(x, d=0.0):
    try:
        return float(x)
    except (TypeError, ValueError):
        return d


# ---------------- Calculation engine (Python) ----------------
SCEN = {"Conservative": -0.03, "Base": 0.0, "Optimistic": 0.03}


def clean(r):
    ia, ea, ts = num(r.get("income_adj"), 100) / 100, num(r.get("exp_adj"), 100) / 100, num(r.get("time_shift"))
    p = dict(
        name=(r.get("name") or "Anonymous").strip()[:80], age=int(num(r.get("age"))),
        income=num(r.get("income")) * ia, expenses=num(r.get("expenses")) * ea,
        debt_payment=num(r.get("debt_payment")), savings=num(r.get("savings")),
        investments=num(r.get("investments")), debt_total=num(r.get("debt_total")),
        liquidity_need=num(r.get("liquidity_need")), risk=r.get("risk", "Medium"),
        return_rate=num(r.get("return_rate"), 6) / 100,
        income_adj=ia * 100, exp_adj=ea * 100, time_shift=ts)
    p["goals"] = [dict(name=(g.get("name") or f"Goal {i+1}")[:60], target=max(num(g.get("target")), 0),
                       years=max(num(g.get("years"), 1) + ts, 0.5), importance=min(max(num(g.get("importance"), 3), 1), 5),
                       saved=max(num(g.get("saved")), 0)) for i, g in enumerate((r.get("goals") or [])[:3])]
    return p


def pmt_needed(target, cur, years, ann):
    n = max(int(round(years * 12)), 1)
    r = (1 + ann) ** (1 / 12) - 1
    factor = ((1 + r) ** n - 1) / r if r else n
    return max(target - cur * (1 + r) ** n, 0) / factor


def evaluate(p, im=1.0, em=1.0, ra=0.0):
    surplus = p["income"] * im - p["expenses"] * em - p["debt_payment"]
    avail = max(surplus * 0.9, 0)  # keep a 10% buffer
    goals = []
    for g in p["goals"]:
        req = pmt_needed(g["target"], g["saved"], g["years"], max(p["return_rate"] + ra, 0))
        aff = 1 - min(req / avail, 1) if avail > 0 else 0
        urg = 1 / (1 + g["years"] / 3)
        score = 0.4 * g["importance"] / 5 + 0.3 * urg + 0.3 * aff
        goals.append({**g, "required": req, "score": round(score, 3)})
    goals.sort(key=lambda g: -g["score"])
    rem = avail
    for i, g in enumerate(goals):
        alloc = min(g["required"], rem)
        rem -= alloc
        g.update(rank=i + 1, funded=alloc, funded_pct=(alloc / g["required"] if g["required"] > 0 else 1))
    tr = sum(g["required"] for g in goals)
    feas = (sum(g["funded"] for g in goals) / tr) if tr > 0 else 1
    return dict(surplus=surplus, avail=avail, goals=goals, total_required=tr, feas=feas)


def analyze(p):
    base = evaluate(p)
    scen = []
    for n, a in SCEN.items():
        e = evaluate(p, ra=a)
        scen.append(dict(name=n, ret=p["return_rate"] + a, required=e["total_required"], available=e["avail"], feasibility=e["feas"]))
    stress = []
    for n, im, em, ra in [("Income -20%", .8, 1, 0), ("Expenses +15%", 1, 1.15, 0),
                          ("Income -20% & Expenses +15%", .8, 1.15, 0), ("Returns -3 pts", 1, 1, -.03)]:
        e = evaluate(p, im, em, ra)
        stress.append(dict(name=n, surplus=e["surplus"], feasibility=e["feas"],
                           runway=(p["savings"] / -e["surplus"] if e["surplus"] < 0 else None)))
    inc, exp = p["income"], p["expenses"]
    emo = p["savings"] / exp if exp else 0
    w = []
    if base["surplus"] < 0: w.append(("danger", f"Negative cash flow: spending and debt payments exceed income by {vnd(-base['surplus'])} per month."))
    dsr = p["debt_payment"] / inc if inc else 0
    if dsr > .4: w.append(("danger", f"Debt payments take {dsr:.0%} of income (above the 40% limit)."))
    elif dsr > .3: w.append(("warn", f"Debt payments take {dsr:.0%} of income (above the 30% comfort level)."))
    if inc and p["debt_total"] / (inc * 12) > 3: w.append(("warn", "Total debt is more than 3x annual income."))
    if emo < 3: w.append(("danger", f"Emergency savings cover only {emo:.1f} months of expenses (target: 6)."))
    elif emo < 6: w.append(("warn", f"Emergency savings cover {emo:.1f} months of expenses (target: 6)."))
    if p["savings"] < p["liquidity_need"]: w.append(("warn", f"Liquid savings are below your {vnd(p['liquidity_need'])} liquidity need."))
    if inc and base["surplus"] / inc < .1: w.append(("warn", f"Savings rate is {max(base['surplus'], 0) / inc:.0%} of income (below 10%)."))
    if base["feas"] < .5: w.append(("danger", f"Only {base['feas']:.0%} of required goal contributions are affordable."))
    elif base["feas"] < .8: w.append(("warn", f"Only {base['feas']:.0%} of required goal contributions are affordable."))
    for g in base["goals"]:
        if g["funded_pct"] < .5: w.append(("warn", f"Goal '{g['name']}' is under-funded ({g['funded_pct']:.0%}). Consider a longer timeline or smaller target."))
    for s in stress:
        if s["surplus"] < 0: w.append(("warn", f"Stress test '{s['name']}': cash flow turns negative ({vnd(s['surplus'])}/month)."))
    assumptions = [f"Expected annual return: {p['return_rate']:.1%} (risk profile: {p['risk']}); not guaranteed.",
                   "Only 90% of monthly surplus is allocated to goals (10% buffer).",
                   "Priority score = 40% importance + 30% urgency + 30% affordability.",
                   "Emergency fund target = 6 months of expenses; funds already saved for goals are separate from savings.",
                   f"What-if adjustments: income {p['income_adj']:.0f}%, expenses {p['exp_adj']:.0f}%, timeline shift {p['time_shift']:+.1f} years."]
    return dict(profile=p, base=base, scenarios=scen, stress=stress, warnings=[dict(level=l, text=t) for l, t in w],
                assumptions=assumptions, emergency_months=emo, savings_rate=(base["surplus"] / inc if inc else 0))


@app.post("/api/analyze")
def api_analyze():
    return jsonify(analyze(clean(request.json or {})))


# ---------------- AI (Google AI Studio / Gemini) ----------------
@app.post("/api/ai")
def api_ai():
    d = request.json or {}
    key = (d.get("api_key") or os.environ.get("GEMINI_API_KEY", "")).strip()
    if not key:
        return jsonify(error="Please enter your Google AI Studio API key."), 400
    a = analyze(clean(d.get("profile") or {}))
    facts = json.dumps({"profile": {k: v for k, v in a["profile"].items() if k != "goals"}, "goals_ranked": a["base"]["goals"],
                        "monthly_surplus": a["base"]["surplus"], "emergency_months": a["emergency_months"],
                        "scenarios": a["scenarios"], "stress_tests": a["stress"], "warnings": a["warnings"]}, ensure_ascii=False)
    prompt = f"""Role: You are a careful personal-finance planning assistant for a bank in Vietnam.
Context: A customer entered a financial profile and 3 goals. All amounts are in VND. A Python engine already computed the numbers below.
Data (use ONLY this): {facts}
Task: Write personalized advice that explains which of the 3 goals should be prioritized and why, the trade-offs between goals, and 3-5 concrete next steps.
Format (English, Markdown, max 450 words): ## Financial snapshot, ## Goal priority & why, ## Trade-offs, ## Next steps, ## Assumptions to confirm.
Constraints: Do not recommend specific securities, funds or products. Do not guarantee returns. Do not invent numbers; if information is missing say "not enough information". Format money as e.g. 1,500,000 VND. End with: "Educational planning only, not investment advice." """
    try:
        r = requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{d.get('model') or DEFAULT_MODEL}:generateContent",
                          headers={"x-goog-api-key": key, "Content-Type": "application/json"},
                          json={"contents": [{"parts": [{"text": prompt}]}]}, timeout=90)
        j = r.json()
        if r.status_code != 200:
            return jsonify(error=f"Gemini API error: {j.get('error', {}).get('message', r.status_code)}"), 502
        return jsonify(text=j["candidates"][0]["content"]["parts"][0]["text"])
    except Exception as e:
        return jsonify(error=f"Could not reach Gemini: {e}"), 502


# ---------------- Excel export + storage ----------------
HEAD = PatternFill("solid", fgColor="14213D")


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    for r in rows: ws.append(r)
    for c in ws[1]: c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(max(len(str(c.value or "")) for c in col) + 3, 70)
    return ws


def client_book(profile_raw, ai_text="", created=""):
    a = analyze(clean(profile_raw)); p = a["profile"]; wb = Workbook(); wb.remove(wb.active)
    sheet(wb, "Profile", [["Field", "Value"], ["Name", p["name"]], ["Age", p["age"]], ["Monthly income (VND)", p["income"]],
          ["Monthly expenses (VND)", p["expenses"]], ["Monthly debt payment (VND)", p["debt_payment"]], ["Savings (VND)", p["savings"]],
          ["Investments (VND)", p["investments"]], ["Total debt (VND)", p["debt_total"]], ["Liquidity need (VND)", p["liquidity_need"]],
          ["Risk tolerance", p["risk"]], ["Expected return", p["return_rate"]], ["Monthly surplus (VND)", a["base"]["surplus"]],
          ["Emergency cover (months)", round(a["emergency_months"], 1)], ["Overall feasibility", a["base"]["feas"]], ["Created", created]])
    sheet(wb, "Goal Priority", [["Rank", "Goal", "Target (VND)", "Years", "Importance", "Saved (VND)", "Required / month", "Funded / month", "Funded %", "Score"]] +
          [[g["rank"], g["name"], g["target"], g["years"], g["importance"], g["saved"], round(g["required"]), round(g["funded"]), g["funded_pct"], g["score"]] for g in a["base"]["goals"]])
    sheet(wb, "Scenarios", [["Scenario", "Return", "Total required / month", "Available / month", "Feasibility"]] +
          [[s["name"], s["ret"], round(s["required"]), round(s["available"]), s["feasibility"]] for s in a["scenarios"]])
    sheet(wb, "Stress Test", [["Test", "Monthly surplus (VND)", "Feasibility", "Runway (months)"]] +
          [[s["name"], round(s["surplus"]), s["feasibility"], round(s["runway"], 1) if s["runway"] else "-"] for s in a["stress"]])
    sheet(wb, "Warnings", [["Level", "Message"]] + [[w["level"], w["text"]] for w in a["warnings"]])
    sheet(wb, "Assumptions", [["Assumption"]] + [[x] for x in a["assumptions"]])
    sheet(wb, "AI Advice", [["AI advice (educational, not investment advice)"]] + [[ln] for ln in (ai_text or "Not generated.").split("\n")])
    return wb


def send_wb(wb, name):
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=name, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/export")
def api_export():
    d = request.json or {}
    return send_wb(client_book(d.get("profile") or {}, d.get("ai_text", ""), dt.datetime.now().isoformat(timespec="seconds")), "financial_plan.xlsx")


@app.post("/api/save")
def api_save():
    d = request.json or {}
    p = clean(d.get("profile") or {})
    with db() as c:
        cur = c.execute("insert into clients(created,name,payload) values(?,?,?)",
                        (dt.datetime.now().isoformat(timespec="seconds"), p["name"], json.dumps(d, ensure_ascii=False)))
    return jsonify(id=cur.lastrowid)


@app.post("/api/admin/export")
def api_admin_export():
    if (request.json or {}).get("password") != ADMIN_PW:
        return jsonify(error="Wrong admin password."), 403
    wb = Workbook(); wb.remove(wb.active)
    cl, gl = [["ID", "Created", "Name", "Age", "Income", "Expenses", "Debt payment", "Savings", "Investments", "Total debt", "Risk", "Surplus", "Emergency months", "Feasibility", "Top priority goal", "Warnings", "AI advice saved"]], \
             [["Client ID", "Name", "Rank", "Goal", "Target", "Years", "Required / month", "Funded %"]]
    for r in db().execute("select * from clients order by id"):
        d = json.loads(r["payload"]); a = analyze(clean(d.get("profile") or {})); p = a["profile"]; g = a["base"]["goals"]
        cl.append([r["id"], r["created"], r["name"], p["age"], p["income"], p["expenses"], p["debt_payment"], p["savings"], p["investments"], p["debt_total"],
                   p["risk"], round(a["base"]["surplus"]), round(a["emergency_months"], 1), a["base"]["feas"], g[0]["name"] if g else "-", len(a["warnings"]), "Yes" if d.get("ai_text") else "No"])
        gl += [[r["id"], r["name"], x["rank"], x["name"], x["target"], x["years"], round(x["required"]), x["funded_pct"]] for x in g]
    sheet(wb, "All Clients", cl); sheet(wb, "All Goals", gl)
    return send_wb(wb, "all_clients.xlsx")


@app.get("/healthz")
def health():
    return "ok"


@app.get("/")
def index():
    return Response(HTML, mimetype="text/html")


HTML = r"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Goal-based Personal Finance Planning Engine</title>
<link href="https://fonts.googleapis.com/css2?family=Manrope:wght@500;700;800&family=Source+Sans+3:wght@400;600&display=swap" rel="stylesheet">
<script src="https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"></script>
<style>
:root{--ink:#14213d;--paper:#f3f5f9;--card:#fff;--teal:#0b6e6e;--teal2:#0f8f8f;--amber:#92400e;--ambg:#fef3c7;--red:#991b1b;--redbg:#fee2e2;--line:#d5dbe6;--mut:#4a5568}
*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:16px/1.5 'Source Sans 3',system-ui,sans-serif}
h1,h2,h3{font-family:Manrope,system-ui,sans-serif;margin:0}
header{background:var(--ink);color:#fff;padding:28px 24px}header h1{font-size:clamp(1.5rem,3vw,2.3rem);font-weight:800}header p{margin:6px 0 0;color:#cbd5e1;max-width:70ch}
.wrap{display:grid;grid-template-columns:minmax(300px,420px) 1fr;gap:20px;padding:20px;max-width:1500px;margin:auto}
@media(max-width:980px){.wrap{grid-template-columns:1fr}}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:18px;margin-bottom:16px}
.card h2{font-size:1.1rem;margin-bottom:10px}label{display:block;font-weight:600;font-size:.9rem;margin-top:10px}
input,select{width:100%;padding:9px 10px;border:1px solid #94a3b8;border-radius:6px;font:inherit;color:var(--ink);background:#fff}
input:focus,select:focus,button:focus-visible{outline:3px solid #5eead4;outline-offset:1px}input[type=range]{padding:0}
.hint{font-size:.8rem;color:var(--teal);min-height:1em}.row{display:grid;grid-template-columns:1fr 1fr;gap:10px}
.goal{border-left:4px solid var(--teal);padding-left:12px;margin-top:14px}
button{font:600 .95rem Manrope,sans-serif;border:0;border-radius:8px;padding:10px 16px;cursor:pointer;background:var(--teal);color:#fff}button:hover{background:var(--teal2)}
button.sec{background:#fff;color:var(--ink);border:2px solid var(--ink)}button.sec:hover{background:#e8edf5}button:disabled{opacity:.5}
.bar{display:flex;flex-wrap:wrap;gap:10px;margin:8px 0}.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin-bottom:16px}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px}.kpi b{display:block;font:800 1.35rem Manrope;margin-top:2px}.kpi span{font-size:.85rem;color:var(--mut)}
table{width:100%;border-collapse:collapse;font-size:.92rem}th,td{padding:8px;border-bottom:1px solid var(--line);text-align:right}th:nth-child(2),td:nth-child(2),th:first-child,td:first-child{text-align:left}th{background:#e8edf5}
.scroll{overflow-x:auto}.charts{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:16px}.charts .card{margin:0}
.w{padding:10px 12px;border-radius:8px;margin:6px 0;font-weight:600}.w.danger{background:var(--redbg);color:var(--red);border-left:5px solid var(--red)}.w.warn{background:var(--ambg);color:var(--amber);border-left:5px solid var(--amber)}.w.ok{background:#dcfce7;color:#14532d;border-left:5px solid #15803d}
#ai{background:#f8fafc;border:1px solid var(--line);border-radius:8px;padding:14px;margin-top:10px}#ai h2{font-size:1rem;margin:12px 0 4px}#ai ul{margin:4px 0;padding-left:20px}
.msg{font-weight:600;margin-top:8px}.err{color:var(--red)}footer{text-align:center;color:var(--mut);padding:20px;font-size:.85rem}
</style></head><body>
<header><h1>Goal-based Personal Finance Planning Engine</h1><p>Enter your finances and three goals. The engine simulates scenarios, stress-tests your plan, flags risks and uses AI to explain which goal to fund first.</p></header>
<div class="wrap"><section>
<div class="card"><h2>1. Your financial profile</h2>
<label>Full name<input id="name" value="Nguyen Van A"></label>
<div class="row"><label>Age<input id="age" type="number" value="32"></label><label>Risk tolerance<select id="risk"><option>Low</option><option selected>Medium</option><option>High</option></select></label></div>
<label>Monthly income (VND)<input class="m" id="income" type="number" value="35000000"></label><div class="hint"></div>
<label>Monthly living expenses (VND)<input class="m" id="expenses" type="number" value="16000000"></label><div class="hint"></div>
<label>Monthly debt payments (VND)<input class="m" id="debt_payment" type="number" value="4000000"></label><div class="hint"></div>
<label>Total outstanding debt (VND)<input class="m" id="debt_total" type="number" value="150000000"></label><div class="hint"></div>
<label>Cash &amp; savings (VND)<input class="m" id="savings" type="number" value="70000000"></label><div class="hint"></div>
<label>Investments (VND)<input class="m" id="investments" type="number" value="40000000"></label><div class="hint"></div>
<label>Liquidity needed within 12 months (VND)<input class="m" id="liquidity_need" type="number" value="30000000"></label><div class="hint"></div></div>
<div class="card"><h2>2. Your three goals</h2><div id="goals"></div></div>
<div class="card"><h2>3. What-if assumptions</h2>
<label>Expected annual return: <b id="v_return_rate">6</b>%<input id="return_rate" type="range" min="1" max="14" step="0.5" value="6"></label>
<label>Income: <b id="v_income_adj">100</b>% of entered<input id="income_adj" type="range" min="60" max="140" step="5" value="100"></label>
<label>Expenses: <b id="v_exp_adj">100</b>% of entered<input id="exp_adj" type="range" min="70" max="140" step="5" value="100"></label>
<label>Move all deadlines by: <b id="v_time_shift">0</b> years<input id="time_shift" type="range" min="-3" max="5" step="1" value="0"></label></div>
<div class="card"><h2>4. AI advisor (Google AI Studio)</h2>
<label>API key<input id="api_key" type="password" placeholder="Paste your Google AI Studio API key" autocomplete="off"></label>
<label>Model<input id="model" value="gemini-3.6-flash"></label>
<label style="display:flex;gap:8px;align-items:flex-start;font-weight:600"><input id="confirm" type="checkbox" style="width:auto;margin-top:5px">I have reviewed the assumptions shown in the results and confirm them.</label>
<div class="bar"><button id="bAI">Get AI advice</button><button class="sec" id="bSave">Save my plan</button><button class="sec" id="bX">Export to Excel</button></div><div class="msg" id="msg"></div></div>
<div class="card"><h2>Administrator</h2><p style="margin:0 0 8px;color:var(--mut);font-size:.9rem">Download every saved customer in one Excel file.</p><button class="sec" id="bAdmin">Export all customers</button></div>
</section><main>
<div class="kpis" id="kpis"></div>
<div class="card"><h2>Goal priority</h2><div class="scroll"><table id="tbl"></table></div></div>
<div class="card"><h2>Financial warnings</h2><div id="warn"></div></div>
<div class="charts"><div class="card"><h2>Required vs. funded per month</h2><canvas id="c1"></canvas></div><div class="card"><h2>Scenario comparison</h2><canvas id="c2"></canvas></div><div class="card"><h2>Stress test: share of goals still affordable</h2><canvas id="c3"></canvas></div></div>
<div class="card" style="margin-top:16px"><h2>Stress test details</h2><div class="scroll"><table id="st"></table></div></div>
<div class="card"><h2>Assumptions used</h2><ul id="asm"></ul></div>
<div class="card"><h2>AI advice</h2><div id="ai">Click "Get AI advice" after confirming the assumptions.</div></div>
</main></div><footer>Educational planning tool. It does not recommend specific securities and does not guarantee returns.</footer>
<script>
const $=id=>document.getElementById(id),F=n=>Math.round(n).toLocaleString('en-US')+' VND',P=n=>Math.round(n*100)+'%';
const GD=[['Home down payment',600000000,8,4,50000000],['Buy a car',350000000,4,3,30000000],['Family trip abroad',80000000,2,2,10000000]];
$('goals').innerHTML=GD.map((g,i)=>`<div class="goal"><h3>Goal ${i+1}</h3><label>Name<input id="gn${i}" value="${g[0]}"></label>
<label>Target amount (VND)<input class="m" id="gt${i}" type="number" value="${g[1]}"></label><div class="hint"></div>
<div class="row"><label>Years to reach<input id="gy${i}" type="number" step="0.5" value="${g[2]}"></label><label>Importance (1-5)<input id="gi${i}" type="number" min="1" max="5" value="${g[3]}"></label></div>
<label>Already saved for it (VND)<input class="m" id="gs${i}" type="number" value="${g[4]}"></label><div class="hint"></div></div>`).join('');
const num=id=>+$(id).value||0;
function collect(){const o={};['name','risk'].forEach(k=>o[k]=$(k).value);['age','income','expenses','debt_payment','debt_total','savings','investments','liquidity_need','return_rate','income_adj','exp_adj','time_shift'].forEach(k=>o[k]=num(k));
o.goals=[0,1,2].map(i=>({name:$('gn'+i).value,target:num('gt'+i),years:num('gy'+i),importance:num('gi'+i),saved:num('gs'+i)}));return o}
document.querySelectorAll('.m').forEach(e=>{const h=e.parentElement.nextElementSibling,f=()=>h.textContent=e.value?F(e.value):'';e.addEventListener('input',f);f()});
['return_rate','income_adj','exp_adj','time_shift'].forEach(k=>$(k).addEventListener('input',()=>$('v_'+k).textContent=$(k).value));
const post=(u,b)=>fetch(u,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(b)});
let A=null,AI='',ch={},t;
function chart(id,cfg){ch[id]&&ch[id].destroy();ch[id]=new Chart($(id),cfg)}
function render(a){A=a;const b=a.base;
$('kpis').innerHTML=[['Monthly surplus',F(b.surplus)],['Savings rate',P(a.savings_rate)],['Emergency cover',a.emergency_months.toFixed(1)+' months'],['Goals affordable',P(b.feas)]].map(k=>`<div class="kpi"><span>${k[0]}</span><b>${k[1]}</b></div>`).join('');
$('tbl').innerHTML='<tr><th>Rank</th><th>Goal</th><th>Target</th><th>Years</th><th>Required / month</th><th>Funded / month</th><th>Funded</th><th>Score</th></tr>'+b.goals.map(g=>`<tr><td>#${g.rank}</td><td>${g.name}</td><td>${F(g.target)}</td><td>${g.years}</td><td>${F(g.required)}</td><td>${F(g.funded)}</td><td>${P(g.funded_pct)}</td><td>${g.score}</td></tr>`).join('');
$('warn').innerHTML=a.warnings.length?a.warnings.map(w=>`<div class="w ${w.level}">${w.text}</div>`).join(''):'<div class="w ok">No major financial warnings detected.</div>';
$('st').innerHTML='<tr><th>Test</th><th>Monthly surplus</th><th>Goals affordable</th><th>Savings runway</th></tr>'+a.stress.map(s=>`<tr><td>${s.name}</td><td>${F(s.surplus)}</td><td>${P(s.feasibility)}</td><td>${s.runway?s.runway.toFixed(1)+' months':'Not needed'}</td></tr>`).join('');
$('asm').innerHTML=a.assumptions.map(x=>`<li>${x}</li>`).join('');
const tick={callback:v=>(v/1e6)+'M'},ttl=c=>c.dataset.label+': '+F(c.parsed.y);
chart('c1',{type:'bar',data:{labels:b.goals.map(g=>g.name),datasets:[{label:'Required',data:b.goals.map(g=>g.required),backgroundColor:'#14213d'},{label:'Funded',data:b.goals.map(g=>g.funded),backgroundColor:'#0f8f8f'}]},options:{scales:{y:{ticks:tick}},plugins:{tooltip:{callbacks:{label:ttl}}}}});
chart('c2',{data:{labels:a.scenarios.map(s=>s.name+' ('+(s.ret*100).toFixed(1)+'%)'),datasets:[{type:'bar',label:'Required / month',data:a.scenarios.map(s=>s.required),backgroundColor:'#b45309'},{type:'line',label:'Available / month',data:a.scenarios.map(s=>s.available),borderColor:'#0f8f8f',backgroundColor:'#0f8f8f',borderWidth:3}]},options:{scales:{y:{ticks:tick}},plugins:{tooltip:{callbacks:{label:ttl}}}}});
chart('c3',{type:'bar',data:{labels:['Base'].concat(a.stress.map(s=>s.name)),datasets:[{label:'Goals affordable',data:[b.feas].concat(a.stress.map(s=>s.feasibility)).map(v=>v*100),backgroundColor:['#0f8f8f','#991b1b','#991b1b','#991b1b','#991b1b']}]},options:{indexAxis:'y',scales:{x:{min:0,max:100,ticks:{callback:v=>v+'%'}}},plugins:{legend:{display:false}}}})}
async function run(){const r=await post('/api/analyze',collect());render(await r.json())}
document.querySelectorAll('input,select').forEach(e=>e.addEventListener('input',()=>{clearTimeout(t);t=setTimeout(run,300)}));
const say=(m,e)=>{$('msg').textContent=m;$('msg').className='msg'+(e?' err':'')};
const need=()=>{if(!$('confirm').checked){say('Please confirm the assumptions first.',1);return false}return true};
function md(s){return s.replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/^## (.*)$/gm,'<h2>$1</h2>').replace(/\*\*(.+?)\*\*/g,'<b>$1</b>').replace(/^[*-] (.*)$/gm,'<li>$1</li>').replace(/(<li>.*<\/li>\n?)+/g,m=>'<ul>'+m+'</ul>').replace(/\n(?!<)/g,'<br>')}
async function dl(r,n){if(!r.ok){say((await r.json()).error||'Download failed.',1);return}const u=URL.createObjectURL(await r.blob()),a=document.createElement('a');a.href=u;a.download=n;a.click();URL.revokeObjectURL(u)}
$('bAI').onclick=async()=>{if(!need())return;say('Asking the AI advisor...');$('bAI').disabled=true;
const r=await post('/api/ai',{profile:collect(),api_key:$('api_key').value,model:$('model').value}),j=await r.json();$('bAI').disabled=false;
if(j.error){say(j.error,1);return}AI=j.text;$('ai').innerHTML=md(AI);say('AI advice ready.')};
$('bSave').onclick=async()=>{if(!need())return;const j=await(await post('/api/save',{profile:collect(),ai_text:AI})).json();say('Plan saved (customer ID '+j.id+').')};
$('bX').onclick=async()=>{if(!need())return;await dl(await post('/api/export',{profile:collect(),ai_text:AI}),'financial_plan.xlsx')};
$('bAdmin').onclick=async()=>{const pw=prompt('Admin password');if(pw)await dl(await post('/api/admin/export',{password:pw}),'all_clients.xlsx')};
run();
</script></body></html>"""

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 7860)))
